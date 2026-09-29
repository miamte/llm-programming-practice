from __future__ import annotations

from pathlib import Path
import re

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


def _extract_section(text: str, title: str) -> str:
    pattern = rf"##\s+{re.escape(title)}\n(.*?)(?=\n##\s+|\Z)"
    match = re.search(pattern, text, re.DOTALL)
    return match.group(1).strip() if match else ""


def _bullets_to_lines(section: str) -> list[str]:
    lines: list[str] = []
    for raw in section.splitlines():
        line = raw.strip()
        if line.startswith("- "):
            lines.append(line[2:].strip())
    return lines


def _set_cell_bg(cell, color_hex: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color_hex)
    tc_pr.append(shd)


def _fill_cell(cell, text: str, bold: bool = False, center: bool = False) -> None:
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
    run = paragraph.add_run(text)
    run.bold = bold
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def create_meeting_docx(summary_md_path: Path, output_docx_path: Path) -> None:
    source = summary_md_path.read_text(encoding="utf-8")

    meeting_title = _extract_section(source, "회의 제목").replace("\n", " ").strip()
    main_points = _bullets_to_lines(_extract_section(source, "주요 내용"))
    positions = _bullets_to_lines(_extract_section(source, "참석자별 입장"))
    decisions = _bullets_to_lines(_extract_section(source, "결정 사항"))

    participants = ", ".join([p.split(":", 1)[0].strip() for p in positions if ":" in p])

    purpose = "기타 구매 시 입문자에게 적합한 선택 기준 정리"
    topic = meeting_title or "기타 구매 기준 논의"

    meeting_body_parts: list[str] = []
    if main_points:
        meeting_body_parts.append("[주요 논의]")
        meeting_body_parts.extend([f"- {item}" for item in main_points])
    if positions:
        meeting_body_parts.append("\n[참석자 입장]")
        meeting_body_parts.extend([f"- {item}" for item in positions])
    meeting_body = "\n".join(meeting_body_parts).strip()

    conclusion = "\n".join([f"- {item}" for item in decisions]) if decisions else "최종 합의 없음"

    document = Document()
    heading = document.add_paragraph("회의록")
    heading.runs[0].bold = True
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER

    table = document.add_table(rows=9, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"

    table.cell(0, 0).text = "날짜"
    table.cell(0, 2).text = "시간"
    table.cell(1, 0).text = "장소"
    table.cell(1, 2).text = "작성자"
    table.cell(2, 0).text = "참석자"
    table.cell(3, 0).text = "목적(목표)"
    table.cell(4, 0).text = "주제"
    table.cell(5, 0).text = "회의 내용"
    table.cell(6, 0).text = "결론 및 합의"
    table.cell(7, 0).text = "할 일"
    table.cell(8, 0).text = "참고 사항"

    table.cell(2, 1).merge(table.cell(2, 3))
    table.cell(3, 1).merge(table.cell(3, 3))
    table.cell(4, 1).merge(table.cell(4, 3))
    table.cell(5, 0).merge(table.cell(5, 0))
    table.cell(5, 1).merge(table.cell(5, 3))
    table.cell(6, 1).merge(table.cell(6, 3))
    table.cell(7, 1).merge(table.cell(7, 3))
    table.cell(8, 1).merge(table.cell(8, 3))

    label_rows = [
        (0, 0), (0, 2),
        (1, 0), (1, 2),
        (2, 0),
        (3, 0),
        (4, 0),
        (5, 0),
        (6, 0),
        (7, 0),
        (8, 0),
    ]
    for r, c in label_rows:
        _set_cell_bg(table.cell(r, c), "DDEAD2")

    _fill_cell(table.cell(0, 0), "날짜", bold=True, center=True)
    _fill_cell(table.cell(0, 2), "시간", bold=True, center=True)
    _fill_cell(table.cell(1, 0), "장소", bold=True, center=True)
    _fill_cell(table.cell(1, 2), "작성자", bold=True, center=True)
    _fill_cell(table.cell(2, 0), "참석자", bold=True, center=True)
    _fill_cell(table.cell(3, 0), "목적(목표)", bold=True, center=True)
    _fill_cell(table.cell(4, 0), "주제", bold=True, center=True)
    _fill_cell(table.cell(5, 0), "회의 내용", bold=True, center=True)
    _fill_cell(table.cell(6, 0), "결론 및 합의", bold=True, center=True)
    _fill_cell(table.cell(7, 0), "할 일", bold=True, center=True)
    _fill_cell(table.cell(8, 0), "참고 사항", bold=True, center=True)

    _fill_cell(table.cell(0, 1), "")
    _fill_cell(table.cell(0, 3), "")
    _fill_cell(table.cell(1, 1), "")
    _fill_cell(table.cell(1, 3), "AI")
    _fill_cell(table.cell(2, 1), participants)
    _fill_cell(table.cell(3, 1), purpose)
    _fill_cell(table.cell(4, 1), topic)
    _fill_cell(table.cell(5, 1), meeting_body)
    _fill_cell(table.cell(6, 1), conclusion)
    _fill_cell(table.cell(7, 1), "- 후속 액션 필요 시 별도 등록")
    _fill_cell(table.cell(8, 1), "원문 요약 파일: guitar_summary.md")

    document.save(output_docx_path)


if __name__ == "__main__":
    base = Path(__file__).resolve().parent
    create_meeting_docx(
        summary_md_path=base / "guitar_summary.md",
        output_docx_path=base / "guitar_meeting_minutes_template.docx",
    )
